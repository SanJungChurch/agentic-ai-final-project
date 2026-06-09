from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "email_to_action_workshop_paper.docx"


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    add_abstract(doc)
    add_body(doc)
    add_references(doc)
    add_appendix(doc)
    doc.save(OUT)
    print(OUT)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, color, before, after in [
        ("Heading 1", 14, "2E74B5", 12, 5),
        ("Heading 2", 11.5, "2E74B5", 8, 3),
        ("Heading 3", 10.5, "1F4D78", 6, 2),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run(
        "Email-to-Action Agent: A LangGraph-Based Workflow for Scheduling, "
        "Place Retrieval, GUI Reservation, and Reply Generation"
    )
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Calibri"

    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_after = Pt(2)
    arun = authors.add_run("Dongwook Kim, Jungin Ju\nSoongsil University, Republic of Korea")
    arun.font.size = Pt(10)
    arun.font.name = "Calibri"

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(8)
    nrun = note.add_run("Workshop paper draft, ACM-style structure")
    nrun.italic = True
    nrun.font.size = Pt(9)
    nrun.font.color.rgb = RGBColor(90, 90, 90)


def add_abstract(doc: Document) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run("Abstract")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("2E74B5")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    text = (
        "We present Email-to-Action Agent, a LangGraph-based agentic workflow that converts "
        "unstructured scheduling e-mails into downstream actions. The system extracts meeting "
        "constraints from Gmail or manually provided e-mail text using Gemini, checks calendar "
        "conflicts, selects a feasible meeting time, retrieves place candidates, attempts GUI-based "
        "reservation through a ShowUI-style executor, and generates a reply draft. Unlike conventional "
        "LLM assistants that stop at summarization or recommendation, our prototype connects language "
        "understanding with external tools including Gmail, Google Calendar, place retrieval providers, "
        "and browser-based reservation pages. A local HTML demo and mock reservation page are used to "
        "evaluate end-to-end behavior under controlled conditions, while Naver Map integration illustrates "
        "the practical constraints of real-world web automation."
    )
    r = p.add_run(text)
    r.font.size = Pt(9.5)

    kw = doc.add_paragraph()
    kw.paragraph_format.space_after = Pt(8)
    k = kw.add_run("Keywords: ")
    k.bold = True
    kw.add_run("LLM agents, scheduling, GUI agents, LangGraph, Gmail, calendar, web automation")


def add_body(doc: Document) -> None:
    add_heading(doc, "1. Introduction", 1)
    add_paragraphs(
        doc,
        [
            "Scheduling a meeting from e-mail is a deceptively simple task. A user must read free-form messages, infer candidate times, check personal calendar conflicts, select a suitable time, search for an appropriate place, attempt a reservation, and write a reply. Current LLM-based assistants are effective at summarizing text or recommending options, but they often stop before taking concrete actions in external environments.",
            "This paper describes Email-to-Action Agent, a prototype that treats e-mail scheduling as an end-to-end agentic workflow rather than a single information extraction problem. The system accepts either manually entered e-mail text or Gmail messages from the recent inbox, extracts scheduling constraints with Gemini, checks conflicts against a calendar source, selects a meeting time, retrieves place candidates, invokes a GUI reservation executor, and drafts a reply.",
            "Our goal is not to claim production-ready autonomous booking, but to investigate a realistic system boundary: how far can an LLM-centered agent proceed when it must combine natural language understanding, tool use, and GUI-level actions?",
        ],
    )

    add_heading(doc, "2. System Design", 1)
    add_paragraphs(
        doc,
        [
            "The system is designed as a stateful workflow with explicit nodes for language understanding, scheduling, retrieval, reservation, and response generation. The graph moves through parse_email, llm_extract, rule_fallback, normalize_times, schedule, retrieve_places, reserve_place, and generate_reply nodes.",
        ],
    )
    add_workflow_table(doc)
    add_heading(doc, "2.1 Constraint Extraction", 2)
    add_paragraphs(
        doc,
        [
            "The extraction node receives an e-mail thread and asks Gemini to produce a structured schema containing intent, participants, candidate times, unavailable times, location preference, meeting duration, missing information, and confidence. Because LLM extraction may fail due to quota, invalid JSON, or missing fields, the graph includes a rule-based fallback extractor and a demo schedule fallback for common Korean scheduling examples.",
        ],
    )
    add_heading(doc, "2.2 Scheduling and Place Retrieval", 2)
    add_paragraphs(
        doc,
        [
            "The scheduling node applies hard constraints for valid time ranges, allowed hours, explicit unavailability, and calendar conflicts. Among valid candidates, it uses a heuristic score based on the number of available participants, preferred times, daytime preference, and meeting duration.",
            "The place retrieval node supports a deterministic demo provider, static HTML provider, mock provider, and Kakao Local API provider. The HTML demo uses the demo provider by default so that changing the e-mail location expression, such as from Soongsil University to Gangnam or Hongdae, changes the recommended place candidate.",
        ],
    )
    add_heading(doc, "2.3 GUI Reservation", 2)
    add_paragraphs(
        doc,
        [
            "Reservation is delegated to a ShowUI-style GUI executor. For controlled evaluation, a local HTML reservation page exposes a machine-readable result marker. For external pages such as Naver Map, the executor inspects visible page text and classifies outcomes as reservation success, unavailable state, login or confirmation gate reached, directions page reached, or unknown.",
        ],
    )

    add_heading(doc, "3. Implementation", 1)
    add_paragraphs(
        doc,
        [
            "The prototype is implemented in Python. LangGraph coordinates workflow state across parsing, extraction, normalization, scheduling, retrieval, reservation, and reply generation nodes. Gemini is used for LLM-based extraction, while Pydantic schemas validate structured outputs.",
            "The HTML demo includes optional Google Workspace integration. A user provides a Google OAuth Desktop client JSON file, after which the system can request Gmail read-only and Calendar read-only scopes. The Gmail loader searches recent messages using the query newer_than:3d, reads up to 50 messages, and filters them using schedule-related keywords such as meeting, schedule, appointment, interview, 회의, 면담, and 일정.",
            "The user-facing demonstration is a standalone HTML page served by a lightweight Python HTTP server. It shows the e-mail input, Gmail loading controls, reservation executor mode, place provider, calendar source, Naver booking URL, and a step-by-step process timeline.",
        ],
    )

    add_heading(doc, "4. Results and Evaluation", 1)
    add_paragraphs(
        doc,
        [
            "We evaluate the prototype through unit tests, controlled end-to-end execution, and real-page observation. The current implementation passes 40 unit tests covering extraction utilities, place retrieval, scheduling, and reservation execution. In the controlled local mock setting, the system completes the full workflow: it extracts candidate meeting times, selects a valid time, recommends a place, submits the reservation form, receives a confirmation ID, and generates a reply draft.",
        ],
    )
    add_evaluation_table(doc)
    add_paragraphs(
        doc,
        [
            "The local mock page is the most reliable evaluation environment because it controls the visible reservation slots and exposes an explicit result state. The real Naver Map experiment is less deterministic. ShowUI can visually operate a browser, but commercial websites still observe browser automation signals, login state, IP reputation, iframe navigation, and request patterns. Therefore, the system reports external outcomes using textual heuristics: success markers, unavailable markers, login gates, directions pages, or unknown results.",
        ],
    )

    add_heading(doc, "5. Limitations", 1)
    add_paragraphs(
        doc,
        [
            "The prototype has several limitations. Gmail and Calendar integration require OAuth configuration, enabled Google APIs, and test-user approval when the OAuth consent screen is in testing mode. LLM extraction can miss or incorrectly normalize dates, participants, and locations, requiring fallback logic for stable demonstrations. Commercial reservation pages may block automated browsers or require login, CAPTCHA, or user confirmation. Finally, the current optimizer is heuristic and does not model travel time, multi-person preference weights, or alternative duration search.",
        ],
    )

    add_heading(doc, "6. Conclusion", 1)
    add_paragraphs(
        doc,
        [
            "Email-to-Action Agent demonstrates how an LLM-centered assistant can be extended into a stateful agentic workflow that connects e-mail understanding with calendar reasoning, place retrieval, GUI reservation attempts, and reply generation. The implementation shows promising controlled end-to-end behavior, while real Naver Map experiments highlight the practical boundary between GUI automation and production web services.",
        ],
    )


def add_workflow_table(doc: Document) -> None:
    rows = [
        ("Input", "Manual HTML input or Gmail schedule-related e-mails"),
        ("Extraction", "Gemini extraction with rule/demo fallback"),
        ("Scheduling", "Calendar conflict check and heuristic optimizer"),
        ("Place Retrieval", "Demo, HTML, mock, Kakao provider options"),
        ("Reservation", "ShowUI-style runner or DOM fallback"),
        ("Reply", "Draft response with selected time, place, and result"),
    ]
    add_table(doc, ["Stage", "Role"], rows)


def add_evaluation_table(doc: Document) -> None:
    rows = [
        ("Gmail loader", "Recent schedule e-mails", "Filters recent 3-day Gmail messages"),
        ("Scheduler", "Valid time selection", "Selects conflict-free candidate"),
        ("Place retrieval", "Location sensitivity", "Changes candidates by location hint"),
        ("Mock reservation", "End-to-end success", "Confirms with local marker"),
        ("Naver target", "Real web action", "Reaches login gate or reports unknown"),
        ("Reply generator", "User response", "Produces ready reply draft"),
    ]
    add_table(doc, ["Component", "Evaluation Target", "Observed Result"], rows)


def add_references(doc: Document) -> None:
    add_heading(doc, "References", 1)
    refs = [
        "LangChain. LangGraph Documentation. https://docs.langchain.com/oss/python/langgraph/overview",
        "ShowUI Authors. ShowUI: One Vision-Language-Action Model for GUI Visual Agent. arXiv:2411.17465, 2024.",
        "Zhou et al. WebArena: A Realistic Web Environment for Building Autonomous Agents. arXiv:2307.13854, 2023.",
        "Liu et al. AgentBench: Evaluating LLMs as Agents. arXiv:2308.03688, 2023.",
        "Google for Developers. Gmail API Documentation. https://developers.google.com/workspace/gmail/api",
        "Google for Developers. Google Calendar API Documentation. https://developers.google.com/workspace/calendar/api",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(ref)


def add_appendix(doc: Document) -> None:
    add_heading(doc, "Appendix: Workflow Nodes", 1)
    add_paragraphs(
        doc,
        [
            "The implemented graph contains the following core nodes: parse_email, llm_extract, rule_fallback, normalize_times, schedule, retrieve_places, reserve_place, and generate_reply. The HTML demo additionally includes Gmail loading, Google Calendar loading, demo schedule fallback, e-mail location override, and Naver URL inference.",
        ],
    )


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def add_paragraphs(doc: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table, 9360)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, "F2F4F7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[idx].paragraphs[0].add_run(value)
    doc.add_paragraph()


def set_table_width(table, width_dxa: int) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


if __name__ == "__main__":
    main()
